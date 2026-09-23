from pathlib import Path

from docx import Document


def extract_docx(source: str | Path | bytes) -> str:
    if isinstance(source, bytes):
        import io
        document = Document(io.BytesIO(source))
    else:
        document = Document(str(source))

    sections = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                sections.append(" | ".join(cells))

    return "\n".join(sections).strip()
