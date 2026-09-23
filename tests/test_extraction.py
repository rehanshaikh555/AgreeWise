from pathlib import Path

import pytest

from extraction import (
    extract_and_normalize,
    extract_text_from_file,
)
from extraction.image import is_ocr_available


def test_txt_extraction(tmp_path: Path):

    file = tmp_path / "policy.txt"

    file.write_text(
        "We collect your personal information.\n\n"
        "We use cookies for analytics.",
        encoding="utf-8",
    )

    result = extract_and_normalize(file)

    assert "personal information" in result
    assert "cookies for analytics" in result


def test_txt_bytes():

    result = extract_and_normalize(
        b"We collect your name and email address.",
        "policy.txt",
    )

    assert "name and email address" in result


def test_docx_extraction(tmp_path: Path):

    from docx import Document

    file = tmp_path / "policy.docx"

    document = Document()

    document.add_paragraph(
        "We collect personal information."
    )

    document.add_paragraph(
        "We may share information with service providers."
    )

    document.save(file)

    result = extract_and_normalize(file)

    assert "personal information" in result
    assert "service providers" in result


def test_pdf_extraction(tmp_path: Path):

    import pymupdf

    file = tmp_path / "policy.pdf"

    document = pymupdf.open()

    page = document.new_page()

    page.insert_text(
        (72, 72),
        "We collect your personal information.",
    )

    document.save(file)
    document.close()

    result = extract_and_normalize(file)

    assert "personal information" in result


def test_normalization():

    text = (
        "   We   collect   your   data.   \n"
        "\n"
        "We use cookies.   "
    )

    result = extract_and_normalize(
        text,
        "policy.txt",
    )

    assert result == (
        "We collect your data.\n"
        "We use cookies."
    )


def test_unsupported_file(tmp_path: Path):

    file = tmp_path / "policy.xyz"

    file.write_text(
        "test",
        encoding="utf-8",
    )

    with pytest.raises(ValueError):
        extract_text_from_file(file)


def test_bytes_require_filename():

    with pytest.raises(ValueError):
        extract_text_from_file(
            b"document",
        )


def test_ocr_status():

    assert isinstance(
        is_ocr_available(),
        bool,
    )
